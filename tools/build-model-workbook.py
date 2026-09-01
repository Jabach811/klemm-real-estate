"""Turns the JSON bundle from export-model-tables.mjs into a workbook plus CSVs.

All of the data logic lives in the Node exporter. This only writes files and
makes the columns readable.

Run: python tools/build-model-workbook.py <bundle.json> <out-dir>
"""

import csv
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MONEY = {"ListPrice", "SalePrice", "PriceVsList", "GrossCommission",
         "AgentNetCommission", "SellerConcession", "MarketingSpend",
         "TypicalMonthlyHOA", "MinPrice", "MaxPrice"}
PERCENT = {"CommissionRate", "TypicalDownPaymentPercent"}
COORD = {"Latitude", "Longitude", "MapX", "MapY"}


def write_sheet(ws, columns, rows):
    ws.append(columns)
    for row in rows:
        ws.append(row)

    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="3F3A33")
    for cell in ws[1]:
        cell.font = head
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions

    for i, name in enumerate(columns, start=1):
        letter = get_column_letter(i)
        width = max(len(str(name)) + 4, 11)
        if rows:
            sample = max(len(str(r[i - 1])) for r in rows[:400])
            width = min(58, max(width, sample + 3))
        ws.column_dimensions[letter].width = width

        fmt = None
        if name in MONEY:
            fmt = '"$"#,##0'
        elif name in PERCENT:
            fmt = "0.00%" if name == "CommissionRate" else "0"
        elif name in COORD:
            fmt = "0.0000"
        elif name == "Date":
            fmt = "yyyy-mm-dd"
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt


def main():
    bundle = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    out_dir = Path(sys.argv[2])
    csv_dir = out_dir / "csv"
    out_dir.mkdir(parents=True, exist_ok=True)
    csv_dir.mkdir(exist_ok=True)

    wb = Workbook()
    wb.remove(wb.active)

    for sheet in bundle["sheets"]:
        cols = sheet["columns"]
        rows = sheet["rows"]
        if "Date" in cols:
            at = cols.index("Date")
            for r in rows:
                r[at] = dt.date.fromisoformat(r[at])
        write_sheet(wb.create_sheet(sheet["name"]), cols, rows)

        with open(csv_dir / f"{sheet['name']}.csv", "w", newline="", encoding="utf-8") as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(rows)

    path = out_dir / "klemm-career-model.xlsx"
    wb.save(path)
    print(f"{len(bundle['sheets'])} sheets -> {path}")
    print(f"matching CSVs -> {csv_dir}")


main()
