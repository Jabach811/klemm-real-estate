from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "docs" / "Mountain-House-Website-Handoff-for-Jack-Updated.docx"
QR_IMAGE = ROOT / "docs" / "handoff-assets" / "mountain-house-site-qr.png"
SITE_URL = "https://klemm-real-estate-tracy.jabach0811.chatgpt.site/mountainhousere.html"

NAVY = RGBColor(25, 49, 73)
BLUE = RGBColor(46, 116, 181)
INK = RGBColor(35, 45, 55)
MUTED = RGBColor(91, 102, 112)
GOLD = RGBColor(173, 120, 46)
PALE_BLUE = "EEF5FA"
PALE_GOLD = "FBF6EC"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def set_cell_margins(paragraph, top=90, bottom=90, start=150, end=150):
    p_pr = paragraph._p.get_or_add_pPr()
    mar = OxmlElement("w:ind")
    mar.set(qn("w:left"), str(start))
    mar.set(qn("w:right"), str(end))
    p_pr.append(mar)
    spacing = p_pr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        p_pr.append(spacing)
    spacing.set(qn("w:before"), str(top))
    spacing.set(qn("w:after"), str(bottom))


def add_bottom_border(paragraph, color="D8E1E8", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2E74B5")
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(fonts)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    new_run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    new_run.append(text_elem)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def set_inline_picture_alt_text(run, description):
    doc_pr = run._element.xpath('.//wp:docPr')[0]
    doc_pr.set('descr', description)
    doc_pr.set('title', 'Mountain House website QR code')


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.78)
    section.bottom_margin = Inches(0.70)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    section.header_distance = Inches(0.30)
    section.footer_distance = Inches(0.32)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h1.font.size = Pt(16)
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    h2.font.size = Pt(13)
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    for name in ("Subtitle", "List Bullet"):
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    bullet.font.size = Pt(10.5)
    bullet.font.color.rgb = INK
    bullet.paragraph_format.left_indent = Inches(0.50)
    bullet.paragraph_format.first_line_indent = Inches(-0.25)
    bullet.paragraph_format.space_after = Pt(7)
    bullet.paragraph_format.line_spacing = 1.167

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("KLEMM REAL ESTATE  |  MOUNTAIN HOUSE WEBSITE HANDOFF  |  SEPTEMBER 2026")
    set_run_font(run, size=8.5, color=MUTED, bold=True)


def add_bullet(doc, lead, detail):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.keep_together = True
    lead_run = p.add_run(lead + " - ")
    set_run_font(lead_run, size=10.5, color=NAVY, bold=True)
    detail_run = p.add_run(detail)
    set_run_font(detail_run, size=10.5, color=INK)


def build_document():
    if not QR_IMAGE.exists():
        raise FileNotFoundError(f"QR image missing: {QR_IMAGE}")

    doc = Document()
    style_document(doc)

    masthead = doc.add_paragraph()
    masthead.paragraph_format.space_before = Pt(0)
    masthead.paragraph_format.space_after = Pt(5)
    masthead.paragraph_format.keep_with_next = True
    run = masthead.add_run("KLEMM REAL ESTATE")
    set_run_font(run, size=10, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("Mountain House Website")
    set_run_font(run, size=25, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.paragraph_format.keep_with_next = True
    run = subtitle.add_run("A polished local sales page, ready to share with Mountain House sellers and buyers.")
    set_run_font(run, size=12, color=MUTED)
    add_bottom_border(subtitle)

    live_label = doc.add_paragraph()
    live_label.paragraph_format.space_before = Pt(10)
    live_label.paragraph_format.space_after = Pt(3)
    live_label.paragraph_format.keep_with_next = True
    shade_paragraph(live_label, PALE_BLUE)
    set_cell_margins(live_label)
    run = live_label.add_run("LIVE WEBSITE")
    set_run_font(run, size=9.5, color=NAVY, bold=True)

    link = doc.add_paragraph()
    link.paragraph_format.space_before = Pt(0)
    link.paragraph_format.space_after = Pt(6)
    link.paragraph_format.keep_with_next = True
    shade_paragraph(link, PALE_BLUE)
    set_cell_margins(link, top=40, bottom=110)
    add_hyperlink(link, SITE_URL, SITE_URL)

    qr = doc.add_paragraph()
    qr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    qr.paragraph_format.space_before = Pt(4)
    qr.paragraph_format.space_after = Pt(1)
    qr.paragraph_format.keep_with_next = True
    qr_run = qr.add_run()
    qr_run.add_picture(str(QR_IMAGE), width=Inches(1.38))
    set_inline_picture_alt_text(qr_run, "QR code linking to the public Klemm Real Estate Mountain House website")

    scan_caption = doc.add_paragraph()
    scan_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scan_caption.paragraph_format.space_before = Pt(0)
    scan_caption.paragraph_format.space_after = Pt(8)
    scan_caption.paragraph_format.keep_with_next = True
    run = scan_caption.add_run("Scan to open the Mountain House site")
    set_run_font(run, size=9.5, color=MUTED, italic=True)

    heading = doc.add_paragraph(style="Heading 1")
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.keep_with_next = True
    heading.add_run("What the site does")

    add_bullet(doc, "Tells the Mountain House story", "A city-specific opening uses the planned-village lifestyle, greenbelts, parks, and Bay Area proximity to make the place feel distinct.")
    add_bullet(doc, "Shows village-level knowledge", "Interactive visuals introduce Wicklund, Bethany, Altamont, Questa, Hansen, and Cordes - reinforcing that pricing is not one-size-fits-all.")
    add_bullet(doc, "Presents listings like a personal marketing piece", "Featured homes are shown as Jack's own listings, not a generic national search feed.")
    add_bullet(doc, "Explains Jack's full-service sale process", "The page walks sellers from the first plan through preparation, marketing, negotiation, escrow, and closing.")
    add_bullet(doc, "Makes the next step obvious", "Clear paths let visitors request a valuation, start a home search, or call and text Jack directly.")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.keep_together = True
    shade_paragraph(note, PALE_GOLD)
    set_cell_margins(note, top=110, bottom=110, start=150, end=150)
    lead = note.add_run("ONE SMALL THING TO FIND\n")
    set_run_font(lead, size=11, color=GOLD, bold=True)
    text = note.add_run("Somewhere in the Mountain House page is a little hidden thing with some significance. You will know it when you find it.")
    set_run_font(text, size=11, color=INK)

    doc.core_properties.title = "Mountain House Website Handoff"
    doc.core_properties.subject = "Klemm Real Estate website handoff for Jack Klemm"
    doc.core_properties.author = "Klemm Real Estate"
    doc.core_properties.comments = "Prepared as a website handoff brief."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
